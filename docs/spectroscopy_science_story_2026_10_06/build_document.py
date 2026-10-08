"""Render the Chinese science document and native vector diagrams.

This renderer imports no Atomx physics modules. Its dependencies are installed
in a separate uv environment, not in the physical calculation environment.
"""
from pathlib import Path
import re
import html
import markdown
import cairosvg
from weasyprint import HTML
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
CSS = """
@page { size:A4; margin:18mm 17mm 18mm;
 @bottom-center { content:'Atomx · 原子分子晶体谱学与未来产业    ' counter(page) ' / ' counter(pages);
 font:8pt 'Noto Sans CJK SC'; color:#617587; }
}
body { font-family:'Noto Sans CJK SC',sans-serif; color:#20303e; font-size:10pt; line-height:1.65; }
h1 { font-size:23pt; line-height:1.4; color:#153d57; margin:0 0 9mm; }
h2 { font-size:17pt; color:#153d57; margin-top:8mm; border-bottom:1px solid #d6e1e8; padding-bottom:2mm; }
h3 { font-size:12pt; color:#294e68; margin-top:6mm; }
h1,h2,h3 { break-after:avoid; }
p { margin:2.5mm 0; orphans:3; widows:3; }
a { color:#24648b; text-decoration:none; overflow-wrap:anywhere; }
table { border-collapse:collapse; width:100%; font-size:8.2pt; line-height:1.5; margin:4mm 0; }
thead { display:table-header-group; } tr { break-inside:avoid; }
th { background:#e6eef4; text-align:left; font-weight:600; }
td,th { padding:1.8mm; border:1px solid #cddbe5; overflow-wrap:anywhere; }
pre { white-space:pre-wrap; background:#f2f5f7; padding:3mm; font-size:8pt; line-height:1.5; overflow-wrap:anywhere; }
code { font-family:'DejaVu Sans Mono','Noto Sans CJK SC',monospace; font-size:0.9em; }
img { width:100%; height:auto; display:block; margin:4mm auto; }
.sources table { font-size:8pt; line-height:1.4; }
.sources td,.sources th { padding:1.2mm; }
li { margin:1.2mm 0; }
@media screen { body { max-width:950px; margin:40px auto; padding:0 24px; } }
"""


def svg_start(width, height):
    return f'''<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">
<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8" fill="#59829b"/></marker></defs>
<rect width="100%" height="100%" fill="#ffffff"/>
<g font-family="Noto Sans CJK SC, sans-serif" fill="#20303e">'''


def label(x, y, value, size=23, fill="#20303e", weight="normal"):
    return f'<text x="{x}" y="{y}" font-size="{size}" fill="{fill}" font-weight="{weight}">{html.escape(value)}</text>'


def box(x, y, w, h, title, lines, fill="#edf4f8"):
    out = f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="16" fill="{fill}" stroke="#bdcedb"/>'
    out += label(x+22, y+40, title, 24, "#153d57", "bold")
    for i, line in enumerate(lines):
        out += label(x+22, y+78+i*32, line, 21)
    return out


def arrow(x1,y1,x2,y2):
    return f'<path d="M{x1},{y1} L{x2},{y2}" fill="none" stroke="#59829b" stroke-width="3" marker-end="url(#arrow)"/>'


def diagrams():
    folder = ROOT / "assets"
    folder.mkdir(exist_ok=True)
    out = svg_start(1200, 710)
    out += label(28,40,"共享原子与芯态物理，逐步加入环境、结构与测量",29,"#153d57","bold")
    out += box(25,70,355,175,"可信原子基础 · S01",["能级 / 相关 / Breit / QED","状态、算符、核与误差预算"])
    out += box(422,70,355,175,"宇宙等离子体 · S02–S03",["太阳 Ne VIII → 恒星 Ne IX/X","布居、流动、密度与 UV 场"])
    out += box(820,70,355,175,"实验室能源 · S04",["Ne / W 杂质辐射与诊断","由原子数据进入输运模型"])
    out += arrow(380,154,417,154)+arrow(777,154,815,154)
    out += box(25,292,355,175,"分子与化学键 · S05",["CO / CO₂ 芯激发与衰变","芯电离 ≠ 中性芯激发"],"#eef5ef")
    out += box(422,292,355,175,"局域环境 · S06",["MgO 嵌入、芯身份与零点","长程电场 ≠ 配体杂化"],"#eef5ef")
    out += box(820,292,355,175,"晶体模型 · S07–S08",["NiO：多重态与电荷转移","TiO₂：延展结构与散射"],"#eef5ef")
    out += arrow(203,245,203,287)+arrow(380,379,417,379)+arrow(777,379,815,379)
    out += box(25,514,754,160,"工作材料与产业决策 · S09–S11",["电池混相 → 催化活性重构 → 半导体界面","用联合观测区分解释，独立验证实际价值"],"#fff5e8")
    out += box(820,514,355,160,"探针与时间 · S12",["Ne XFEL 电荷级联","剂量和动态控制贯穿全链"],"#fff5e8")
    out += arrow(997,467,997,509)+arrow(997,494,402,494)+arrow(402,494,402,509)
    out += '</g></svg>'
    (folder / "science_story.svg").write_text(out,encoding="utf-8")
    out = svg_start(1200,390)
    out += label(25,42,"一次有科学判断力的 agent 研究循环",29,"#153d57","bold")
    out += box(25,75,270,180,"问题与解释",["观测与原始来源","至少两个竞争解释","预先声明验证条件"])
    out += box(320,75,270,180,"模型与计算",["选最小判别控制","锁定输入与物理模型","估算资源并执行"])
    out += box(615,75,270,180,"验证与结果",["契约 / 守恒 / 收敛","独立数据与联合观测","保存失败及误差"])
    out += box(910,75,265,180,"下一步决定",["支持 / 排除 / 未分辨","追加控制或新测量","发布可回放证据"])
    out += arrow(295,164,315,164)+arrow(590,164,610,164)+arrow(885,164,905,164)
    out += '<path d="M1040,255 L1040,300 L160,300 L160,260" fill="none" stroke="#59829b" stroke-width="3" marker-end="url(#arrow)"/>'
    out += label(330,340,"新证据更新假设；接口通过、谱学验证与产业有效性分别报告",23)
    out += '</g></svg>'
    (folder / "agent_loop.svg").write_text(out,encoding="utf-8")
    for src in folder.glob("*.svg"):
        cairosvg.svg2png(url=str(src),write_to=str(src.with_suffix(".png")),scale=2)


def configure_font(style,size):
    style.font.name="Noto Sans CJK SC"
    style.font.size=Pt(size)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Noto Sans CJK SC")


def inline(paragraph,text):
    pattern = r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)"
    for part in re.split(pattern,text):
        if not part: continue
        m = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)",part)
        if m:
            element=OxmlElement("w:hyperlink")
            relation=paragraph.part.relate_to(m[2],"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
            element.set(qn("r:id"),relation)
            run=OxmlElement("w:r"); props=OxmlElement("w:rPr")
            color=OxmlElement("w:color"); color.set(qn("w:val"),"24648B"); props.append(color); run.append(props)
            t=OxmlElement("w:t");t.text=m[1];run.append(t);element.append(run);paragraph._p.append(element)
        elif part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold=True
        elif part.startswith("`"):
            paragraph.add_run(part[1:-1])
        else: paragraph.add_run(part)


def docx_document(stem,text):
    doc=Document(); sec=doc.sections[0]
    sec.page_width,sec.page_height=Cm(21),Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(1.8)
    sec.left_margin=sec.right_margin=Cm(1.7)
    for style,size in [("Normal",10),("Title",23),("Heading 1",17),("Heading 2",12)]:configure_font(doc.styles[style],size)
    doc.styles["Normal"].paragraph_format.line_spacing=1.5
    doc.styles["Normal"].paragraph_format.space_after=Pt(6)
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line.startswith("```"):
            block=[];i+=1
            while i<len(lines) and not lines[i].startswith("```"):block.append(lines[i]);i+=1
            p=doc.add_paragraph("\n".join(block));p.paragraph_format.line_spacing=1.15
            for run in p.runs:run.font.size=Pt(8)
        elif line.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].startswith("|"):
                cells=[cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+",cell) for cell in cells):rows.append(cells)
                i+=1
            table=doc.add_table(rows=1,cols=len(rows[0]));table.style="Table Grid"
            for ri,cells in enumerate(rows):
                row=table.rows[0] if ri==0 else table.add_row()
                cant=OxmlElement("w:cantSplit");row._tr.get_or_add_trPr().append(cant)
                if ri==0:
                    repeat=OxmlElement("w:tblHeader");row._tr.get_or_add_trPr().append(repeat)
                for cell,value in zip(row.cells,cells):
                    p=cell.paragraphs[0];inline(p,value);p.paragraph_format.line_spacing=1.25
                    for run in p.runs:run.font.size=Pt(8)
            continue
        elif line.startswith("# "):doc.add_heading(line[2:],0)
        elif line.startswith("## "):doc.add_heading(line[3:],1)
        elif line.startswith("### "):doc.add_heading(line[4:],2)
        elif line.startswith("!["):
            m=re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)",line)
            path=ROOT/m[2]
            if path.suffix==".svg":path=path.with_suffix(".png")
            doc.add_picture(str(path),width=Cm(17.3));doc.add_paragraph(m[1])
        else:
            style=None
            if line.startswith("- "):line=line[2:];style="List Bullet"
            if re.match(r"\d+\. ",line):line=re.sub(r"^\d+\. ","",line);style="List Number"
            p=doc.add_paragraph(style=style);inline(p,line.rstrip())
        i+=1
    footer=sec.footer.paragraphs[0];footer.alignment=1
    footer.add_run("Atomx · 原子分子晶体谱学与未来产业  ")
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)
    doc.save(ROOT/f"{stem}.docx")


def render(stem):
    text=(ROOT/f"{stem}.md").read_text(encoding="utf-8")
    body=markdown.markdown(text,extensions=["tables","fenced_code"])
    source_heading='<h2>12. 原始来源与软件入口</h2>'
    if source_heading in body:
        body=body.replace(source_heading,'<section class="sources">'+source_heading,1)+'</section>'
    title=html.escape(text.splitlines()[0].lstrip("# "))
    full=f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>{title}</title><style>{CSS}</style><body>{body}</body></html>'
    (ROOT/f"{stem}.html").write_text(full,encoding="utf-8")
    HTML(string=full,base_url=str(ROOT)).write_pdf(ROOT/f"{stem}.pdf")
    docx_document(stem,text)


if __name__=="__main__":
    diagrams()
    for stem in ["研究体系与案例设计","故事线摘要"]:render(stem)
    print("Rendered two native diagrams and main/summary HTML, PDF, DOCX.")

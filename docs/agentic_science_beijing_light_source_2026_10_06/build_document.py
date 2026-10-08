"""Render the competition document without importing any Atomx physics modules."""
from pathlib import Path
import html
import re

import markdown
from weasyprint import HTML
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
TEXT = (ROOT / "参赛项目书.md").read_text(encoding="utf-8")
CSS = """
@page { size: A4; margin: 18mm 17mm 18mm;
 @bottom-center { content: 'Atomx · 北京光源与未来产业    ' counter(page) ' / ' counter(pages);
 font: 8pt 'Noto Sans CJK SC'; color: #617587; }
}
body { font-family:'Noto Sans CJK SC',sans-serif; color:#20303e; font-size:10pt; line-height:1.7; }
h1 { font-size:24pt; line-height:1.4; color:#153d57; margin:0 0 12mm; }
h2 { font-size:17pt; color:#153d57; margin-top:9mm; border-bottom:1px solid #d6e1e8; padding-bottom:2mm; }
h3 { font-size:12pt; color:#294e68; margin-top:6mm; }
h1,h2,h3 { break-after:avoid; }
h2:last-of-type { break-before:page; }
p { margin:2.5mm 0; orphans:3; widows:3; }
a { color:#24648b; text-decoration:none; overflow-wrap:anywhere; }
table { border-collapse:collapse; width:100%; font-size:8.2pt; line-height:1.55; margin:4mm 0; }
thead { display:table-header-group; } tr { break-inside:avoid; }
th { background:#e6eef4; text-align:left; font-weight:600; }
td,th { padding:2mm; border:1px solid #cddbe5; overflow-wrap:anywhere; }
pre { white-space:pre-wrap; background:#f2f5f7; padding:3mm; font-size:8.2pt; line-height:1.6; overflow-wrap:anywhere; }
code { font-family:'DejaVu Sans Mono','Noto Sans CJK SC',monospace; }
img { width:100%; height:auto; display:block; margin:4mm auto; }
li { margin:1.2mm 0; }
@media screen { body { max-width:920px; margin:40px auto; padding:0 24px; } }
"""


def plain(s):
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    return s.replace("**", "").replace("`", "")


def configure_font(style, size):
    style.font.name = "Noto Sans CJK SC"
    style.font.size = Pt(size)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")


def add_inline(p, s):
    for match in re.finditer(r"\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|([^\[]+?)(?=\[|\*\*|$)", s):
        label, url, bold, ordinary = match.groups()
        if url:
            hyperlink = OxmlElement("w:hyperlink")
            rel = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            hyperlink.set(qn("r:id"), rel)
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            color = OxmlElement("w:color"); color.set(qn("w:val"), "24648B"); properties.append(color)
            run.append(properties)
            t = OxmlElement("w:t"); t.text = label; run.append(t)
            hyperlink.append(run); p._p.append(hyperlink)
        elif bold:
            p.add_run(bold).bold = True
        elif ordinary:
            p.add_run(ordinary.replace("`", ""))


def render_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.7)
    for name, size in [("Normal",10),("Title",23),("Heading 1",17),("Heading 2",12)]:
        configure_font(doc.styles[name],size)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.5
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    lines = TEXT.splitlines(); i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip(): i += 1; continue
        if line.startswith("```"):
            block=[]; i += 1
            while i<len(lines) and not lines[i].startswith("```"): block.append(lines[i]); i+=1
            p=doc.add_paragraph("\n".join(block)); p.paragraph_format.line_spacing=1.15
            for r in p.runs:r.font.size=Pt(8)
        elif line.startswith("|"):
            rows=[]
            while i<len(lines) and lines[i].startswith("|"):
                cells=[plain(c.strip()) for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+",c) for c in cells):rows.append(cells)
                i+=1
            table=doc.add_table(rows=1,cols=len(rows[0]));table.style="Table Grid"
            for row_index,cells in enumerate(rows):
                row=table.rows[0] if row_index==0 else table.add_row()
                cant=OxmlElement("w:cantSplit");row._tr.get_or_add_trPr().append(cant)
                if row_index==0:
                    repeat=OxmlElement("w:tblHeader");row._tr.get_or_add_trPr().append(repeat)
                for cell,s in zip(row.cells,cells):
                    cell.text=s
                    for p in cell.paragraphs:
                        p.paragraph_format.line_spacing=1.25
                        for run in p.runs:run.font.size=Pt(8);run.bold=row_index==0
            continue
        elif line.startswith("# "):
            doc.add_heading(line[2:],0)
        elif line.startswith("## "):
            doc.add_heading(line[3:],1)
        elif line.startswith("### "):
            doc.add_heading(line[4:],2)
        elif line.startswith("!["):
            m=re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)",line)
            path=ROOT/m[2]
            if path.suffix==".svg": path=ROOT/"assets/agent_workflow.png"
            doc.add_picture(str(path),width=Cm(17.3))
            doc.add_paragraph(m[1])
        else:
            style=None
            if re.match(r"\d+\. ",line):line=re.sub(r"^\d+\. ","",line);style="List Number"
            p=doc.add_paragraph(style=style);add_inline(p,line.rstrip())
        i+=1
    footer=section.footer.paragraphs[0];footer.alignment=1
    footer.add_run("Atomx · 北京光源与未来产业  ")
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)
    doc.save(ROOT/"参赛项目书.docx")


if __name__=="__main__":
    import cairosvg
    cairosvg.svg2png(url=str(ROOT/"assets/agent_workflow.svg"),write_to=str(ROOT/"assets/agent_workflow.png"),scale=2)
    body=markdown.markdown(TEXT,extensions=["tables","fenced_code"])
    full=f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>Atomx 北京光源参赛项目书</title><style>{CSS}</style><body>{body}</body></html>'
    (ROOT/"参赛项目书.html").write_text(full,encoding="utf-8")
    HTML(string=full,base_url=str(ROOT)).write_pdf(ROOT/"参赛项目书.pdf")
    render_docx()
    print("Rendered HTML, PDF, DOCX and workflow image.")
